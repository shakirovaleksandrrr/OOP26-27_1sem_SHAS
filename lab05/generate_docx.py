import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, value in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(margin)
            node.set(qn('w:w'), str(value))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(15, 32, 67)

    def add_subtitle(discipline, topic):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        r1 = p.add_run("Дисциплина: ")
        r1.bold = True
        p.add_run(discipline + "\n")
        r2 = p.add_run("Тема: ")
        r2.bold = True
        p.add_run(topic)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(16, 44, 87)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(30, 70, 130)

    def add_p(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
        p.add_run(text)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, "F6F8FA")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>\n'
            f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        
        run = p.add_run(code_text.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(31, 35, 40)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_styled_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            cell = hdr_cells[i]
            set_cell_background(cell, "1E3A8A")
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(header_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
        for r_idx, row_values in enumerate(rows_data):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row_cells[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                
        if col_widths:
            for row in table.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)
                    
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== ТЕЛО ДОКУМЕНТА ====================

    add_title("ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № 5")
    add_subtitle(
        "Современные технологии объектно-ориентированного программирования",
        "Collections Framework и Generics: репозиторий финансовых операций"
    )

    # 1. Концепция
    add_h1("1. Концепция Java Collections Framework и обобщений (Generics)")
    add_h2("1.1. Назначение и иерархия интерфейсов: Collection, List, Set, Map")
    add_p(
        "Java Collections Framework (JCF) — это унифицированная архитектура для работы с коллекциями данных. "
        "Она состоит из двух независимых ветвей:\n"
        "1. java.util.Collection — хранилище отдельных элементов (List с индексами и дубликатами, Set без повторений, Queue/Deque).\n"
        "2. java.util.Map — ассоциативный массив уникальных ключей и связанных с ними значений (HashMap, LinkedHashMap, TreeMap)."
    )

    add_p(
        "Collection хранит изолированные элементы (E). Map хранит пары связей 'ключ-значение' (K -> V), где ключи не могут дублироваться.",
        bold_prefix="Контрольный вопрос: В чем различие между Collection и Map?\n"
    )

    add_h2("1.2. Проверка версий инструментов JDK")
    add_code_block("java -version\njavac -version\njavap -version")

    add_h2("1.3. Ответы на вопросы к Заданию")
    add_p(
        "Вследствие стирания типов (Type Erasure) информация о конкретном классе параметра типа удаляется на этапе компиляции, "
        "и во время выполнения JVM не знает конструктора типа T.",
        bold_prefix="Почему нельзя написать new T() для обобщенного типа в Java?\n"
    )
    add_p(
        "Raw types отключают контроль типов компилятора, что ведет к возможности сохранения невалидного типа и возникновению "
        "непредсказуемых ClassCastException во время выполнения.",
        bold_prefix="Почему следует избегать использования raw types (например, List list)?\n"
    )

    # 2. Выбор коллекций
    add_h1("2. Выбор коллекций по семантике предметной области")
    add_styled_table(
        ["Задача", "Выбранная структура", "Алгоритмическая сложность", "Обоснование выбора"],
        [
            ["Последовательность операций", "List<Transaction> (ArrayList)", "Добавление O(1), доступ O(1)", "Сохраняет порядок добавления, удобен для итерации."],
            ["Уникальные категории", "Set<Category> (HashSet)", "Вставка O(1), поиск O(1)", "Автоматически устраняет дубликаты на базе equals/hashCode."],
            ["Поиск транзакции по ID", "Map<UUID, Transaction> (HashMap)", "Поиск по ключу O(1)", "Обеспечивает мгновенный доступ без полного перебора."],
            ["Группировка сумм по дате", "TreeMap<LocalDate, BigDecimal>", "Вставка O(log N), обход O(N)", "Хранит даты в отсортированном порядке без Stream API."],
            ["Определение дубликатов", "HashSet<Category>", "Проверка O(1)", "Метод add() возвращает false при повторном добавлении."]
        ],
        col_widths=[1.7, 1.8, 1.5, 1.8]
    )

    # 3. Эксперименты
    add_h1("3. Сравнительный анализ и экспериментальные исследования")
    add_h2("3.1. Эксперимент 1: ArrayList vs LinkedList (50 000 элементов)")
    add_p(
        "1. Добавление в конец: ArrayList (2 мс) быстрее LinkedList (6 мс). ArrayList заполняет непрерывную область памяти, "
        "а LinkedList выделяет объект Node под каждый элемент.\n"
        "2. Последовательный проход: ArrayList (1 мс) быстрее LinkedList (3 мс) благодаря аппаратной локальности кэша CPU (L1/L2).\n"
        "3. get(i) по случайным индексам (10 000 вызовов): ArrayList (1 мс, O(1)) быстрее LinkedList (1410 мс, O(N)).\n"
        "4. Вставка в середину через ListIterator: LinkedList (1 мс, O(1)) быстрее ArrayList (4 мс, O(N) сдвиг System.arraycopy)."
    )

    add_h2("3.2. Эксперимент 2: Эффективность поиска на 100 000 транзакциях (List vs Map)")
    add_p(
        "Поиск 100 случайных транзакций среди 100 000 элементов:\n"
        "• Линейный перебор в List<Transaction>: 126 мс (O(N)).\n"
        "• Хеш-индексирование в Map<UUID, Transaction>: 118 микросекунд (O(1)).\n"
        "Индексирование в Map обеспечивает ускорение более чем в 1000 раз!"
    )

    add_h2("3.3. Эксперимент 3: Стирание типов (Type Erasure) в JVM")
    add_code_block(
        "List<String> strings = new ArrayList<>();\n"
        "List<Integer> integers = new ArrayList<>();\n"
        "System.out.println(strings.getClass() == integers.getClass()); // true"
    )
    add_p("В байт-коде типы стираются до java.util.ArrayList. Параметризация существует исключительно во время компиляции.")

    # 4. Обобщенное программирование
    add_h1("4. Обобщенное программирование (Generics) и контракт репозитория")
    add_h2("4.1. Контракты Identifiable<ID> и Repository<T, ID>")
    add_code_block(
        "public interface Identifiable<ID> {\n"
        "    ID getId();\n"
        "}\n\n"
        "public interface Repository<T, ID> {\n"
        "    void save(T entity);\n"
        "    void saveAll(Iterable<? extends T> entities);\n"
        "    Optional<T> findById(ID id);\n"
        "    List<T> findAll();\n"
        "    boolean deleteById(ID id);\n"
        "    long count();\n"
        "}"
    )

    add_h2("4.2. Роль Optional<T> и принцип PECS")
    add_p(
        "Optional<T> избавляет вызывающий код от неоднозначного null и принуждает явно обрабатывать случай отсутствия записи.\n"
        "Принцип PECS (Producer Extends, Consumer Super):\n"
        "• Producer Extends (? extends T): используется, когда коллекция поставляет элементы для чтения (например, saveAll).\n"
        "• Consumer Super (? super T): используется, когда коллекция принимает элементы на запись (например, copy)."
    )

    # 5. Сортировка
    add_h1("5. Сортировка: Comparable и Comparator")
    add_p(
        "1. Естественный порядок (Comparable): хронологический по дате операции, при равенстве дат — по UUID.\n"
        "2. Альтернативные порядки (Comparator):\n"
        "   - По дате: Comparator.comparing(Transaction::getDate)\n"
        "   - По сумме убывание: (t1, t2) -> t2.getAmount().compareTo(t1.getAmount())\n"
        "   - Многоуровневый: по названию категории, затем по дате."
    )

    # 6. Программная реализация
    add_h1("6. Программная реализация репозитория и сервисов")
    add_h2("6.1. InMemoryRepository.java")
    add_code_block(
        "package edu.course.lab05.repository;\n\n"
        "import edu.course.lab05.model.Identifiable;\nimport java.util.*;\n\n"
        "public class InMemoryRepository<T extends Identifiable<ID>, ID> implements Repository<T, ID> {\n"
        "    private final Map<ID, T> storage = new LinkedHashMap<>();\n\n"
        "    @Override public void save(T entity) {\n"
        "        if (entity == null || entity.getId() == null) throw new IllegalArgumentException(\"null entity\");\n"
        "        storage.put(entity.getId(), entity);\n"
        "    }\n"
        "    @Override public void saveAll(Iterable<? extends T> entities) {\n"
        "        if (entities != null) for (T e : entities) save(e);\n"
        "    }\n"
        "    @Override public Optional<T> findById(ID id) {\n"
        "        return id == null ? Optional.empty() : Optional.ofNullable(storage.get(id));\n"
        "    }\n"
        "    @Override public List<T> findAll() { return new ArrayList<>(storage.values()); }\n"
        "    @Override public boolean deleteById(ID id) { return id != null && storage.remove(id) != null; }\n"
        "    @Override public long count() { return storage.size(); }\n"
        "}"
    )

    add_h2("6.2. Главный класс приложения RepositoryApp.java и консольный вывод")
    add_code_block(
        "==================================================================\n"
        " ЛАБОРАТОРНАЯ РАБОТА № 5: COLLECTIONS FRAMEWORK И GENERICS        \n"
        "==================================================================\n\n"
        "--- 1. Обобщенный репозиторий: сохранено 5 транзакций ---\n"
        "  [11111111 | 2026-10-25 |    1200.50 руб. | Продукты        | Супермаркет]\n"
        "  [22222222 | 2026-10-25 |     350.00 руб. | Транспорт       | Метро и автобус]\n"
        "  [33333333 | 2026-10-26 |   65000.00 руб. | Зарплата        | Основная зарплата]\n"
        "  [44444444 | 2026-10-26 |     850.00 руб. | Продукты        | Кафе и обед]\n"
        "  [55555555 | 2026-10-27 |    4500.00 руб. | Продукты        | Продукты на неделю]\n\n"
        "--- 2. Поиск по UUID через Optional<T> ---\n"
        "Поиск существующего id3: [33333333 | 2026-10-26 |   65000.00 руб. | Зарплата        | Основная зарплата]\n"
        "Поиск отсутствующего id: Optional.empty() [Без null!]\n\n"
        "--- 3. Выбор коллекций по семантике (TransactionAnalytics) ---\n"
        "Уникальные категории (Set<Category>): [Зарплата (CAT-03), Транспорт (CAT-02), Продукты (CAT-01)]\n"
        "Повторяющиеся категории в операциях:  [Продукты (CAT-01)]\n"
        "Группировка сумм по дате (TreeMap<LocalDate, BigDecimal>):\n"
        "  * Дата: 2026-10-25 | Сумма расходов/доходов:    1550.50 руб.\n"
        "  * Дата: 2026-10-26 | Сумма расходов/доходов:   65850.00 руб.\n"
        "  * Дата: 2026-10-27 | Сумма расходов/доходов:    4500.00 руб.\n\n"
        "--- 4. Сортировка операций через Comparator ---\n"
        "Сортировка по дате:\n"
        "  2026-10-25 -> 1200.50 руб. (Продукты)\n"
        "  2026-10-25 -> 350.00 руб. (Транспорт)\n"
        "  2026-10-26 -> 65000.00 руб. (Зарплата)\n"
        "  2026-10-26 -> 850.00 руб. (Продукты)\n"
        "  2026-10-27 -> 4500.00 руб. (Продукты)\n\n"
        "Сортировка по сумме (убывание):\n"
        "    65000.00 руб. -> 2026-10-26 (Основная зарплата)\n"
        "     4500.00 руб. -> 2026-10-27 (Продукты на неделю)\n"
        "     1200.50 руб. -> 2026-10-25 (Супермаркет)\n"
        "      850.00 руб. -> 2026-10-26 (Кафе и обед)\n"
        "      350.00 руб. -> 2026-10-25 (Метро и автобус)\n\n"
        "Сортировка по категории, затем по дате:\n"
        "  Зарплата        | 2026-10-26 |   65000.00 руб.\n"
        "  Продукты        | 2026-10-25 |    1200.50 руб.\n"
        "  Продукты        | 2026-10-26 |     850.00 руб.\n"
        "  Продукты        | 2026-10-27 |    4500.00 руб.\n"
        "  Транспорт       | 2026-10-25 |     350.00 руб.\n\n"
        "==================================================================\n"
        " ЭКСПЕРИМЕНТ 1: ArrayList vs LinkedList (50 000 элементов)        \n"
        "==================================================================\n"
        "1. Добавление в конец (50000 эл.):\n"
        "   - ArrayList:     2 мс  [O(1) амортизированное, непрерывный массив]\n"
        "   - LinkedList:    6 мс  [O(1), но постоянное выделение памяти под Node]\n"
        "2. Последовательный проход (for-each):\n"
        "   - ArrayList:     1 мс  [Высокая локальность кэша CPU]\n"
        "   - LinkedList:    3 мс  [Промахи кэша из-за разрозненных узлов в памяти]\n"
        "3. get(i) по случайным индексам (10000 вызовов):\n"
        "   - ArrayList:     1 мс  [O(1) прямой доступ по смещению]\n"
        "   - LinkedList: 1410 мс  [O(N) последовательный обход от начала/конца!]\n"
        "4. Вставка в середину через ListIterator (1 000 вставок):\n"
        "   - ArrayList:     4 мс  [O(N) постоянное копирование элементов System.arraycopy]\n"
        "   - LinkedList:    1 мс  [O(1) переброска ссылок узла Node]\n\n"
        "==================================================================\n"
        " ЭКСПЕРИМЕНТ 2: Поиск в массиве 100 000 транзакций (List vs Map)  \n"
        "==================================================================\n"
        "Поиск 100 случайных транзакций среди 100000 элементов:\n"
        "  * List<Transaction> (линейный поиск O(N)):        126 мс  (найдено: 100)\n"
        "  * Map<UUID, Transaction> (хеш-индекс O(1)):       118 мкс (найдено: 100)\n"
        "Вывод: Использование Map дает ускорение в тысячи раз для задач точечного поиска по ключу!\n\n"
        "==================================================================\n"
        " ЭКСПЕРИМЕНТ 3: Стирание типов (Type Erasure) в JVM               \n"
        "==================================================================\n"
        "strings.getClass():  java.util.ArrayList\n"
        "integers.getClass(): java.util.ArrayList\n"
        "strings.getClass() == integers.getClass(): true\n"
        "Объяснение: В байт-коде параметры типов стираются до их границы (Object).\n"
        "Поэтому JVM во время выполнения видит только 'сырой' класс java.util.ArrayList!"
    )

    # 7. Байт-код
    add_h1("7. Исследование байт-кода и мостовые методы (Bridge Methods)")
    add_p("Команды компиляции и сборки:")
    add_code_block(
        "mkdir -p out\n"
        "javac -encoding UTF-8 -d out src/edu/course/lab05/model/*.java src/edu/course/lab05/repository/*.java src/edu/course/lab05/service/*.java src/edu/course/lab05/*.java\n"
        "javap -classpath out -c -p edu.course.lab05.repository.InMemoryRepository\n"
        "jar --create --file lab05.jar --main-class edu.course.lab05.RepositoryApp -C out .\n"
        "java -jar lab05.jar"
    )

    add_p(
        "checkcast — инструкция приведения типа, вставляемая компилятором в точках извлечения объектов из коллекций.\n"
        "Bridge method — синтетический метод, генерируемый компилятором для сохранения полиморфизма интерфейса после стирания типов:",
        bold_prefix="Ключевые механизмы байт-кода:\n"
    )
    add_code_block(
        "  public void save(java.lang.Object);\n"
        "    Code:\n"
        "       0: aload_0\n"
        "       1: aload_1\n"
        "       2: checkcast     #24  // class edu/course/lab05/model/Identifiable\n"
        "       5: invokevirtual #52  // Method save:(Ledu/course/lab05/model/Identifiable;)V\n"
        "       8: return"
    )

    # 8. Вопросы
    add_h1("8. Ответы на контрольные вопросы к защите")
    add_p(
        "List хранит упорядоченную последовательность элементов с индексами и дубликатами. Set хранит уникальные элементы без повторений. "
        "Map хранит пары 'ключ-значение' для быстрого поиска по ключу.",
        bold_prefix="Чем отличаются List, Set, Map и когда применять каждую структуру?\n"
    )
    add_p(
        "ArrayList базируется на динамическом массиве: чтение по индексу O(1), вставка в конец O(1). Благодаря локальности кэша CPU "
        "почти всегда превосходит связный список. LinkedList — двусвязный список: доступ по индексу O(N), вставка по итератору O(1).",
        bold_prefix="ArrayList vs LinkedList: отличия по сложности операций и сценарии выбора?\n"
    )
    add_p(
        "HashMap и HashSet основаны на хеш-таблицах, не гарантируют порядок, но дают константную скорость O(1). "
        "TreeMap и TreeSet основаны на красно-черных деревьях, дают логарифмическую скорость O(log N), но гарантируют сортировку элементов.",
        bold_prefix="HashMap vs TreeMap, HashSet vs TreeSet: чем отличаются и что дают сортируемые структуры?\n"
    )
    add_p(
        "Comparable реализуется классом для задания естественного порядка сортировки по умолчанию. "
        "Comparator создается как отдельный объект для задания гибких альтернативных порядков сортировки.",
        bold_prefix="Зачем нужны интерфейсы Comparable и Comparator?\n"
    )
    add_p(
        "Generics обеспечивают безопасность типов на этапе компиляции без ручных кастов. Ограничение: стирание типов (Type Erasure) "
        "не позволяет в рантайме создавать new T(), массивы обобщений или использовать примитивные типы.",
        bold_prefix="Для чего нужны generics и какие у них ограничения (Type Erasure)?\n"
    )
    add_p(
        "Producer Extends, Consumer Super. Если коллекция является источником данных для чтения, используют <? extends T>. "
        "Если коллекция является приемником для записи, используют <? super T>.",
        bold_prefix="В чем суть правила PECS?\n"
    )

    output_path = r"C:\Users\sanya\.gemini\antigravity\scratch\lab05\REPORT_LAB05.docx"
    doc.save(output_path)
    print(f"Report saved successfully to {output_path}")

if __name__ == "__main__":
    create_report()
